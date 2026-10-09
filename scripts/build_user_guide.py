#!/usr/bin/env python3
"""
Build every edition of the user guide, and the help inside BillDesk, from ONE source file.

    python scripts/build_user_guide.py

Reads    Docs/BillDesk_Native_User_Guide.src.md      (chapters = '## ', topics = '### ' each followed by a <!-- meta --> line)
         Docs/user-guide-native/img/*.png            (screenshots, made by scripts/user_guide/capture_screenshots.py)
Writes   Docs/BillDesk_Native_User_Guide.md          contents, numbering, anchors, "see also", A-Z index
         Docs/BillDesk_Native_User_Guide.docx / .pdf / .ipynb
         app/assets/help/guide.json + img/           what the Help screen (F9) searches and shows

Meta line, after the heading:   <!-- id: bill-date | screen: New Bill | keywords: a, b | related: x, y | context: Page A; Page B -->
Links between topics are written [text](#topic-id). The build stops with a message if an id, a link, a related topic or an image is missing.
Needs python-docx, reportlab and Pillow.
"""
from __future__ import annotations

import base64
import json
import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "Docs"
SRC = DOCS / "BillDesk_Native_User_Guide.src.md"
STEM = DOCS / "BillDesk_Native_User_Guide"
HELP_DIR = ROOT / "app" / "assets" / "help"
TITLE = "BillDesk — User Guide and Help"


# ====================================================================================================== model
@dataclass
class Topic:
    id: str
    title: str
    chapter: str = ""
    number: str = ""
    screen: str = ""
    keywords: List[str] = field(default_factory=list)
    related: List[str] = field(default_factory=list)
    context: List[str] = field(default_factory=list)
    lines: List[str] = field(default_factory=list)
    blocks: List[dict] = field(default_factory=list)


@dataclass
class Chapter:
    id: str
    title: str
    number: int = 0
    topics: List[Topic] = field(default_factory=list)


@dataclass
class Guide:
    title: str
    front_lines: List[str]
    chapters: List[Chapter]

    def topics(self) -> List[Topic]:
        return [t for c in self.chapters for t in c.topics]

    def by_id(self) -> Dict[str, Topic]:
        return {t.id: t for t in self.topics()}


def parse_meta(line: str) -> Dict[str, str]:
    m = re.match(r"^\s*<!--\s*(.*?)\s*-->\s*$", line)
    if not m:
        return {}
    meta = {}
    for part in m.group(1).split("|"):
        if ":" in part:
            k, v = part.split(":", 1)
            meta[k.strip().lower()] = v.strip()
    return meta


def parse_source(text: str) -> Guide:
    lines = text.splitlines()
    title = TITLE
    front: List[str] = []
    chapters: List[Chapter] = []
    cur_ch: Optional[Chapter] = None
    cur_t: Optional[Topic] = None
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# ") and not chapters:
            title = ln[2:].strip()
        elif ln.startswith("## "):
            meta = parse_meta(lines[i + 1]) if i + 1 < len(lines) else {}
            cid = meta.get("chapter") or re.sub(r"\W+", "-", ln[3:].lower()).strip("-")
            cur_ch = Chapter(cid, meta.get("title") or re.sub(r"^Chapter \d+\.\s*", "", ln[3:].strip()), len(chapters) + 1)
            chapters.append(cur_ch)
            cur_t = None
            i += 1 if meta else 0
        elif ln.startswith("### ") and cur_ch is not None:
            meta = parse_meta(lines[i + 1]) if i + 1 < len(lines) else {}
            if "id" not in meta:
                raise SystemExit(f"topic '{ln[4:]}' has no <!-- id: ... --> line")
            cur_t = Topic(id=meta["id"], title=ln[4:].strip(), chapter=cur_ch.id, number=f"{cur_ch.number}.{len(cur_ch.topics) + 1}",
                          screen=meta.get("screen", ""), keywords=[k.strip() for k in meta.get("keywords", "").split(",") if k.strip()],
                          related=[r.strip() for r in meta.get("related", "").split(",") if r.strip()],
                          context=[c.strip() for c in meta.get("context", "").split(";") if c.strip()])
            cur_ch.topics.append(cur_t)
            i += 1
        elif re.match(r"^-{3,}$", ln.strip()):
            pass
        elif ln.strip().startswith("<!--"):
            pass
        elif cur_t is not None:
            cur_t.lines.append(ln)
        elif cur_ch is None:
            front.append(ln)
        i += 1
    guide = Guide(title, front, chapters)
    for t in guide.topics():
        t.blocks = parse_blocks(t.lines)
    return guide


def parse_blocks(lines: List[str]) -> List[dict]:
    blocks: List[dict] = []
    i = 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append({"t": "code", "text": "\n".join(code)})
            continue
        if ln.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append({"t": "table", "rows": rows})
            continue
        if ln.startswith(">"):
            q = []
            while i < len(lines) and lines[i].startswith(">"):
                q.append(lines[i].lstrip("> ").rstrip())
                i += 1
            blocks.append({"t": "quote", "text": " ".join(x for x in q if x)})
            continue
        m = re.match(r"^\s*!\[(.*?)\]\((.*?)\)\s*$", ln)
        if m:
            blocks.append({"t": "image", "alt": m.group(1), "src": m.group(2)})
            i += 1
            continue
        m = re.match(r"^(\s*)([*-]|\d+\.)\s+(.*)$", ln)
        if m:
            ordered = m.group(2)[0].isdigit()
            items = []
            while i < len(lines):
                mm = re.match(r"^(\s*)([*-]|\d+\.)\s+(.*)$", lines[i])
                if not mm or mm.group(2)[0].isdigit() != ordered:
                    break
                text = mm.group(3)
                i += 1
                while i < len(lines) and lines[i].strip() and lines[i].startswith("   ") and not re.match(r"^\s*([*-]|\d+\.)\s", lines[i]):
                    text += " " + lines[i].strip()
                    i += 1
                items.append(text)
            blocks.append({"t": "ol" if ordered else "ul", "items": items})
            continue
        para = [ln.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(```|\||>|\s*[*-]\s|\s*\d+\.\s|\s*!\[)", lines[i]):
            para.append(lines[i].strip())
            i += 1
        blocks.append({"t": "p", "text": " ".join(para)})
    return blocks


# ------------------------------------------------------------------------------------------------------ text helpers
INLINE = re.compile(r"(\[[^\]]+\]\(#[A-Za-z0-9_-]+\)|\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*)")
LINK = re.compile(r"\[([^\]]+)\]\(#([A-Za-z0-9_-]+)\)")


def inline(text: str) -> List[tuple]:
    """(kind, text, extra) runs: plain '' / b / i / code / link(extra = topic id)."""
    out = []
    for part in INLINE.split(text):
        if not part:
            continue
        lm = LINK.fullmatch(part)
        if lm:
            out.append(("link", lm.group(1), lm.group(2)))
        elif part.startswith("**") and part.endswith("**"):
            out.append(("b", part[2:-2], None))
        elif part.startswith("`") and part.endswith("`"):
            out.append(("code", part[1:-1], None))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append(("i", part[1:-1], None))
        else:
            out.append(("", part, None))
    return out


def plain(text: str) -> str:
    return "".join(t for _k, t, _e in inline(text))


def block_text(b: dict) -> str:
    t = b["t"]
    if t in ("p", "quote", "code"):
        return plain(b["text"]) if t != "code" else b["text"]
    if t in ("ul", "ol"):
        return "\n".join(plain(x) for x in b["items"])
    if t == "table":
        return "\n".join(" | ".join(plain(c) for c in r) for r in b["rows"])
    if t == "image":
        return b.get("alt", "")
    return ""


# ====================================================================================================== validation
def validate(guide: Guide) -> None:
    ids = [t.id for t in guide.topics()]
    dup = {i for i in ids if ids.count(i) > 1}
    problems = [f"duplicate topic id: {d}" for d in dup]
    known = set(ids)
    for t in guide.topics():
        for r in t.related:
            if r not in known:
                problems.append(f"{t.id}: related topic '{r}' does not exist")
        text = "\n".join(t.lines)
        for _label, target in LINK.findall(text):
            if target not in known:
                problems.append(f"{t.id}: link to unknown topic '#{target}'")
        for b in t.blocks:
            if b["t"] == "image" and not (DOCS / b["src"]).exists():
                problems.append(f"{t.id}: image missing: {b['src']}")
    for line in guide.front_lines:
        for _l, target in LINK.findall(line):
            if target not in known:
                problems.append(f"front matter: link to unknown topic '#{target}'")
    if problems:
        raise SystemExit("The guide has problems:\n  " + "\n  ".join(problems))


# ====================================================================================================== derived data
def contents_lines(guide: Guide) -> List[str]:
    out = []
    for c in guide.chapters:
        out.append(f"- **Chapter {c.number}. {c.title}**")
        for t in c.topics:
            out.append(f"  - [{t.number} {t.title}](#{t.id})")
    return out


def keyword_index(guide: Guide) -> Dict[str, List[Topic]]:
    idx: Dict[str, List[Topic]] = {}
    for t in guide.topics():
        for k in {*t.keywords, t.title.lower()}:
            idx.setdefault(k.lower(), []).append(t)
    return dict(sorted(idx.items(), key=lambda kv: kv[0]))


def see_also(guide: Guide, t: Topic) -> str:
    by = guide.by_id()
    links = [f"[{by[r].title}](#{r})" for r in t.related if r in by]
    return ("**See also:** " + " · ".join(links)) if links else ""


def published_markdown(guide: Guide) -> str:
    L: List[str] = [f"# {guide.title}", ""]
    L += [ln for ln in guide.front_lines if not ln.startswith("# ") and not ln.strip().startswith("<!--")]
    L += ["", "## Contents", ""] + contents_lines(guide) + ["", "- [Index of words](#index)", "", "---", ""]
    for c in guide.chapters:
        L += [f'<a id="chapter-{c.id}"></a>', f"## Chapter {c.number}. {c.title}", ""]
        for t in c.topics:
            L += [f'<a id="{t.id}"></a>', f"### {t.number} {t.title}", ""]
            if t.screen:
                L += [f"*Screen: {t.screen}*", ""]
            L += [ln for ln in t.lines]
            sa = see_also(guide, t)
            if sa:
                L += ["", sa]
            L += ["", "---", ""]
    L += ['<a id="index"></a>', "## Index of words", "", "*Look up a word, then follow the link.*", ""]
    letter = ""
    for k, topics in keyword_index(guide).items():
        if k[:1].upper() != letter:
            letter = k[:1].upper()
            L += ["", f"**{letter}**", ""]
        L.append(f"- {k} — " + ", ".join(f"[{t.number}](#{t.id})" for t in topics[:6]))
    return "\n".join(L).rstrip() + "\n"


# ====================================================================================================== help data (in-program)
def build_help(guide: Guide) -> dict:
    from PIL import Image
    img_out = HELP_DIR / "img"
    img_out.mkdir(parents=True, exist_ok=True)
    for old in img_out.glob("*.png"):
        old.unlink()
    by = guide.by_id()
    ch_title = {c.id: c.title for c in guide.chapters}
    topics = {}
    for t in guide.topics():
        blocks = []
        for b in t.blocks:
            b = dict(b)
            if b["t"] == "image":
                name = Path(b["src"]).name
                src = DOCS / b["src"]
                im = Image.open(src)
                if im.width > 900:
                    im = im.resize((900, int(im.height * 900 / im.width)), Image.LANCZOS)
                im.save(img_out / name, optimize=True)
                b["src"] = name
            blocks.append(b)
        first_p = next((plain(b["text"]) for b in t.blocks if b["t"] == "p"), "") or             next((plain(b["items"][0]) for b in t.blocks if b["t"] in ("ol", "ul") and b["items"]), "")
        first_ol = next((b["items"] for b in t.blocks if b["t"] == "ol"), [])
        heads = [plain(x) for b in t.blocks if b["t"] == "table" for x in [r[0] for r in b["rows"][1:]] if x]
        topics[t.id] = {
            "id": t.id, "title": t.title, "number": t.number, "chapter": t.chapter, "chapter_title": ch_title[t.chapter], "screen": t.screen,
            "keywords": t.keywords, "related": [r for r in t.related if r in by], "summary": first_p, "steps": first_ol, "headings": heads[:40],
            "text": "\n".join(block_text(b) for b in t.blocks), "blocks": blocks,
        }
    intents = []
    faq = []
    for t in guide.topics():
        for b in t.blocks:
            if b["t"] != "table" or not b["rows"]:
                continue
            head = [h.lower() for h in b["rows"][0]]
            for row in b["rows"][1:]:
                links = LINK.findall(" ".join(row))
                if head[:1] == ["i want to…"] and links:
                    intents.append({"phrase": plain(row[0]), "topic": links[0][1]})
                elif head[:1] == ["question"] and len(row) >= 3:
                    faq.append({"q": plain(row[0]), "a": plain(row[1]), "topic": (LINK.findall(row[2]) or [(None, t.id)])[0][1]})
                elif head[:1] == ["what you see"] and len(row) >= 3:
                    target = (LINK.findall(row[2]) or [(None, "messages")])[0][1]
                    faq.append({"q": plain(row[0]), "a": plain(row[1]) + " " + plain(row[2]), "topic": target})
    context = {}
    for t in guide.topics():
        for page in t.context:
            context[page] = t.id
    data = {"title": guide.title, "chapters": [{"id": c.id, "title": f"{c.number}. {c.title}", "topics": [t.id for t in c.topics]} for c in guide.chapters],
            "topics": topics, "intents": intents, "faq": faq, "context": context}
    HELP_DIR.mkdir(parents=True, exist_ok=True)
    (HELP_DIR / "guide.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    return data


# ====================================================================================================== Word
def build_docx(guide: Guide):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from PIL import Image

    d = Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for n, sz in (("Heading 1", 22), ("Heading 2", 16), ("Heading 3", 13)):
        d.styles[n].font.name = "Calibri"
        d.styles[n].font.size = Pt(sz)
        d.styles[n].font.color.rgb = RGBColor(0x1E, 0x29, 0x6B)
    width_cm = 17.0
    bm_id = [100]

    def bookmark(par, name):
        bm_id[0] += 1
        s = OxmlElement("w:bookmarkStart")
        s.set(qn("w:id"), str(bm_id[0]))
        s.set(qn("w:name"), name)
        e = OxmlElement("w:bookmarkEnd")
        e.set(qn("w:id"), str(bm_id[0]))
        par._p.insert(0, s)
        par._p.append(e)

    def internal_link(par, text, anchor, size=None, bold=False):
        h = OxmlElement("w:hyperlink")
        h.set(qn("w:anchor"), anchor)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        col = OxmlElement("w:color")
        col.set(qn("w:val"), "4F46E5")
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(col)
        rpr.append(u)
        if bold:
            rpr.append(OxmlElement("w:b"))
        if size:
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), str(int(size * 2)))
            rpr.append(sz)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        h.append(r)
        par._p.append(h)

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), hexcolor)
        tcPr.append(sh)

    def runs(par, text, size=None, bold=False):
        for kind, t, extra in inline(text):
            if kind == "link":
                internal_link(par, t, extra, size, bold)
                continue
            r = par.add_run(t)
            r.bold = bold or kind == "b"
            r.italic = kind == "i"
            if size:
                r.font.size = Pt(size)
            if kind == "code":
                r.font.name = "Consolas"
                r.font.size = Pt((size or 10.5) - 0.5)

    def render(blocks):
        for b in blocks:
            t = b["t"]
            if t == "p":
                runs(d.add_paragraph(), b["text"])
            elif t in ("ul", "ol"):
                for n, item in enumerate(b["items"], start=1):
                    p = d.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.9)
                    p.paragraph_format.first_line_indent = Cm(-0.55)
                    p.paragraph_format.space_after = Pt(2)
                    p.add_run(("•" if t == "ul" else f"{n}.") + "  ")
                    runs(p, item)
            elif t == "quote":
                p = d.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                pPr = p._p.get_or_add_pPr()
                bd = OxmlElement("w:pBdr")
                lf = OxmlElement("w:left")
                for k, v in (("val", "single"), ("sz", "18"), ("space", "8"), ("color", "5B54D6")):
                    lf.set(qn("w:" + k), v)
                bd.append(lf)
                pPr.append(bd)
                runs(p, b["text"])
            elif t == "code":
                p = d.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                r = p.add_run(b["text"])
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)
            elif t == "table":
                rows = b["rows"]
                ncol = max(len(r) for r in rows)
                tb = d.add_table(rows=len(rows), cols=ncol)
                tb.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        cell = tb.cell(ri, ci)
                        cell.text = ""
                        runs(cell.paragraphs[0], row[ci] if ci < len(row) else "", size=9.5, bold=(ri == 0))
                        if ri == 0:
                            shade(cell, "E0E7FF")
                d.add_paragraph()
            elif t == "image":
                p = DOCS / b["src"]
                w_px = Image.open(p).size[0]
                w = width_cm if w_px >= 1000 else min(width_cm, w_px / 96 * 2.54 * 0.9)
                par = d.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.add_run().add_picture(str(p), width=Cm(w))
                cap = d.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(b["alt"])
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    d.add_heading(guide.title, 0)
    render(parse_blocks([ln for ln in guide.front_lines if not ln.startswith("# ") and not ln.strip().startswith("<!--")]))
    d.add_page_break()
    d.add_heading("Contents", 1)
    for c in guide.chapters:
        p = d.add_paragraph()
        internal_link(p, f"Chapter {c.number}. {c.title}", f"chapter-{c.id}", bold=True)
        for t in c.topics:
            q = d.add_paragraph()
            q.paragraph_format.left_indent = Cm(0.8)
            q.paragraph_format.space_after = Pt(0)
            internal_link(q, f"{t.number}  {t.title}", t.id, size=10)
    p = d.add_paragraph()
    internal_link(p, "Index of words", "index", bold=True)
    note = d.add_paragraph()
    r = note.add_run("Tip: press Ctrl+F in Word and type what you are looking for, or open the Navigation Pane (View → Navigation Pane) to jump between chapters and topics.")
    r.italic = True
    by = guide.by_id()
    for c in guide.chapters:
        d.add_page_break()
        hp = d.add_heading(f"Chapter {c.number}. {c.title}", 1)
        bookmark(hp, f"chapter-{c.id}")
        for t in c.topics:
            tp = d.add_heading(f"{t.number}  {t.title}", 2)
            bookmark(tp, t.id)
            if t.screen:
                s = d.add_paragraph()
                r = s.add_run(f"Screen: {t.screen}")
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            render(t.blocks)
            links = [r for r in t.related if r in by]
            if links:
                p = d.add_paragraph()
                p.add_run("See also: ").bold = True
                for n, r in enumerate(links):
                    if n:
                        p.add_run("  ·  ")
                    internal_link(p, by[r].title, r)
    d.add_page_break()
    ip = d.add_heading("Index of words", 1)
    bookmark(ip, "index")
    letter = ""
    for k, topics in keyword_index(guide).items():
        if k[:1].upper() != letter:
            letter = k[:1].upper()
            lp = d.add_paragraph()
            lp.add_run(letter).bold = True
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(0)
        p.add_run(k + " — ")
        for n, t in enumerate(topics[:6]):
            if n:
                p.add_run(", ")
            internal_link(p, t.number, t.id)
    # page numbers in the footer
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        el = OxmlElement("w:fldChar" if kind else "w:instrText")
        if kind:
            el.set(qn("w:fldCharType"), kind)
        else:
            el.text = txt
            el.set(qn("xml:space"), "preserve")
        run._r.append(el)
    d.core_properties.title = guide.title
    d.save(str(STEM) + ".docx")


# ====================================================================================================== PDF
def build_pdf(guide: Guide):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (BaseDocTemplate, Frame, HRFlowable, Image as RLImage, KeepTogether, PageBreak, PageTemplate, Paragraph,
                                    Preformatted, Spacer, Table, TableStyle)
    from reportlab.platypus.tableofcontents import TableOfContents
    from PIL import Image

    fonts = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    regs = {"UI": "segoeui.ttf", "UI-B": "segoeuib.ttf", "UI-I": "segoeuii.ttf", "UI-BI": "segoeuiz.ttf", "MONO": "consola.ttf"}
    if all((fonts / f).exists() for f in regs.values()):
        for n, f in regs.items():
            pdfmetrics.registerFont(TTFont(n, str(fonts / f)))
        pdfmetrics.registerFontFamily("UI", normal="UI", bold="UI-B", italic="UI-I", boldItalic="UI-BI")
        base, bold, ital, mono = "UI", "UI-B", "UI-I", "MONO"
    else:
        base, bold, ital, mono = "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Courier"
    ink, accent, navy = colors.HexColor("#1e293b"), colors.HexColor("#4f46e5"), colors.HexColor("#1e296b")
    S = {
        "p": ParagraphStyle("p", fontName=base, fontSize=9.5, leading=13.5, textColor=ink, spaceAfter=5),
        "title": ParagraphStyle("title", fontName=bold, fontSize=26, leading=32, textColor=accent, spaceAfter=8, alignment=TA_CENTER),
        "ch": ParagraphStyle("ch", fontName=bold, fontSize=20, leading=24, textColor=navy, spaceBefore=0, spaceAfter=8),
        "tp": ParagraphStyle("tp", fontName=bold, fontSize=13.5, leading=17, textColor=navy, spaceBefore=12, spaceAfter=4),
        "screen": ParagraphStyle("screen", fontName=ital, fontSize=8, leading=10, textColor=colors.HexColor("#64748b"), spaceAfter=4),
        "cap": ParagraphStyle("cap", fontName=ital, fontSize=8, leading=10, textColor=colors.HexColor("#64748b"), alignment=TA_CENTER, spaceAfter=8),
        "cell": ParagraphStyle("cell", fontName=base, fontSize=8.5, leading=11, textColor=ink),
        "code": ParagraphStyle("code", fontName=mono, fontSize=8.5, leading=11, backColor=colors.HexColor("#f1f5f9"), leftIndent=8, borderPadding=5, spaceAfter=8),
        "quote": ParagraphStyle("quote", fontName=base, fontSize=9.5, leading=13.5, textColor=colors.HexColor("#334155"), leftIndent=12, backColor=colors.HexColor("#eef2ff"), borderPadding=6, spaceAfter=8),
        "idx": ParagraphStyle("idx", fontName=base, fontSize=8.5, leading=11, leftIndent=8, spaceAfter=0),
        "letter": ParagraphStyle("letter", fontName=bold, fontSize=11, leading=14, textColor=accent, spaceBefore=6),
    }

    def esc(t):
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def mk(text):
        o = []
        for kind, t, extra in inline(text):
            t = esc(t)
            if kind == "link":
                o.append(f'<a href="#{extra}" color="#4f46e5">{t}</a>')
            else:
                o.append({"b": f"<b>{t}</b>", "i": f"<i>{t}</i>", "code": f'<font name="{mono}" color="#9d174d">{t}</font>'}.get(kind, t))
        return "".join(o)

    class Heading(Paragraph):
        def __init__(self, text, style, level, key, label):
            super().__init__(f'<a name="{key}"/>{esc(text)}', style)
            self.toc_level, self.toc_key, self.toc_label = level, key, text

    class Doc(BaseDocTemplate):
        def afterFlowable(self, fl):
            if isinstance(fl, Heading):
                self.canv.bookmarkPage(fl.toc_key)
                self.canv.addOutlineEntry(fl.toc_label, fl.toc_key, fl.toc_level, closed=fl.toc_level > 0)
                self.notify("TOCEntry", (fl.toc_level, fl.toc_label, self.page, fl.toc_key))

    usable = A4[0] - 4 * cm

    def render(blocks, story):
        for b in blocks:
            t = b["t"]
            if t == "p":
                story.append(Paragraph(mk(b["text"]), S["p"]))
            elif t in ("ul", "ol"):
                for n, item in enumerate(b["items"], start=1):
                    mark = "•" if t == "ul" else f"{n}."
                    story.append(Paragraph(f"{mark}&nbsp;&nbsp;" + mk(item), ParagraphStyle("li", parent=S["p"], leftIndent=16, firstLineIndent=-12, spaceAfter=2)))
            elif t == "quote":
                story.append(Paragraph(mk(b["text"]), S["quote"]))
            elif t == "code":
                story.append(Preformatted(b["text"], S["code"]))
            elif t == "table":
                rows = b["rows"]
                ncol = max(len(r) for r in rows)
                data = [[Paragraph(("<b>%s</b>" % mk(c)) if ri == 0 else mk(c), S["cell"]) for c in (r + [""] * (ncol - len(r)))] for ri, r in enumerate(rows)]
                tb = Table(data, colWidths=[usable / ncol] * ncol, repeatRows=1)
                tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e0e7ff")), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                                        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
                story += [tb, Spacer(1, 8)]
            elif t == "image":
                p = DOCS / b["src"]
                w_px, h_px = Image.open(p).size
                w = usable if w_px >= 1000 else min(usable, w_px * 0.75)
                h = w * h_px / w_px
                if h > 17 * cm:
                    h = 17 * cm
                    w = h * w_px / h_px
                story.append(KeepTogether([RLImage(str(p), width=w, height=h), Paragraph(esc(b["alt"]), S["cap"])]))

    story = [Spacer(1, 3 * cm), Paragraph(esc(guide.title), S["title"])]
    render(parse_blocks([ln for ln in guide.front_lines if not ln.startswith("# ") and not ln.strip().startswith("<!--")]), story)
    story.append(PageBreak())
    story.append(Paragraph("Contents", S["ch"]))
    toc = TableOfContents()
    toc.levelStyles = [ParagraphStyle("t0", fontName=bold, fontSize=10.5, leading=15, spaceBefore=5, leftIndent=0, textColor=navy),
                       ParagraphStyle("t1", fontName=base, fontSize=9, leading=12, leftIndent=16, textColor=ink)]
    story += [toc, Paragraph("<i>Tip: use the bookmarks panel of your PDF reader, or Ctrl+F, to find a topic quickly.</i>", S["screen"])]
    by = guide.by_id()
    for c in guide.chapters:
        story += [PageBreak(), Heading(f"Chapter {c.number}. {c.title}", S["ch"], 0, f"chapter-{c.id}", f"Chapter {c.number}. {c.title}")]
        for t in c.topics:
            story.append(Heading(f"{t.number}  {t.title}", S["tp"], 1, t.id, f"{t.number}  {t.title}"))
            if t.screen:
                story.append(Paragraph(f"Screen: {esc(t.screen)}", S["screen"]))
            render(t.blocks, story)
            links = [r for r in t.related if r in by]
            if links:
                story.append(Paragraph("<b>See also:</b> " + " · ".join(f'<a href="#{r}" color="#4f46e5">{esc(by[r].title)}</a>' for r in links), S["p"]))
            story.append(HRFlowable(width="100%", thickness=0.4, color=colors.HexColor("#e2e8f0"), spaceBefore=4, spaceAfter=2))
    story += [PageBreak(), Heading("Index of words", S["ch"], 0, "index", "Index of words")]
    letter = ""
    for k, topics in keyword_index(guide).items():
        if k[:1].upper() != letter:
            letter = k[:1].upper()
            story.append(Paragraph(letter, S["letter"]))
        story.append(Paragraph(esc(k) + " — " + ", ".join(f'<a href="#{t.id}" color="#4f46e5">{t.number}</a>' for t in topics[:6]), S["idx"]))

    def footer(c, doc):
        c.saveState()
        c.setFont(base, 8)
        c.setFillColor(colors.HexColor("#64748b"))
        c.drawString(2 * cm, 1.1 * cm, guide.title)
        c.drawRightString(A4[0] - 2 * cm, 1.1 * cm, f"Page {doc.page}")
        c.restoreState()

    doc = Doc(str(STEM) + ".pdf", pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm, title=guide.title, author="BillDesk")
    doc.addPageTemplates([PageTemplate(id="main", frames=[Frame(2 * cm, 1.8 * cm, usable, A4[1] - 3.6 * cm, id="f")], onPage=footer)])
    doc.multiBuild(story)


# ====================================================================================================== Jupyter notebook
def build_ipynb(guide: Guide, help_data: dict):
    """One markdown cell per topic (screenshots embedded as attachments) and, at the end, a code cell with the very same search the program's Help uses."""
    cells = []

    def md_cell(text: str):
        att = {}

        def repl(m):
            alt, src = m.group(1), m.group(2)
            p = DOCS / src
            if not p.exists():
                return m.group(0)
            att[p.name] = {"image/png": base64.b64encode(p.read_bytes()).decode("ascii")}
            return f"![{alt}](attachment:{p.name})"
        body = re.sub(r"!\[(.*?)\]\((.*?)\)", repl, text)
        cell = {"cell_type": "markdown", "metadata": {}, "source": body.rstrip().splitlines(keepends=True)}
        if att:
            cell["attachments"] = att
        cells.append(cell)

    front = [ln for ln in guide.front_lines if not ln.startswith("# ") and not ln.strip().startswith("<!--")]
    md_cell(f"# {guide.title}\n\n" + "\n".join(front) + "\n\n**Search this guide:** run the last cell once, then type `ask(\"your question\")` in any code cell (for example `ask(\"how do I return goods\")`).\n\n## Contents\n\n" + "\n".join(contents_lines(guide)) + "\n\n- [Index of words](#index)")
    for c in guide.chapters:
        md_cell(f'<a id="chapter-{c.id}"></a>\n\n## Chapter {c.number}. {c.title}')
        for t in c.topics:
            lines = [f'<a id="{t.id}"></a>', "", f"### {t.number} {t.title}", ""]
            if t.screen:
                lines += [f"*Screen: {t.screen}*", ""]
            lines += t.lines
            sa = see_also(guide, t)
            if sa:
                lines += ["", sa]
            md_cell("\n".join(lines))
    idx = ['<a id="index"></a>', "", "## Index of words", ""]
    for k, topics in keyword_index(guide).items():
        idx.append(f"- {k} — " + ", ".join(f"[{t.number}](#{t.id})" for t in topics[:6]))
    md_cell("\n".join(idx))

    service_src = (ROOT / "app" / "services" / "help_service.py").read_text(encoding="utf-8")
    service_src = re.sub(r"(?m)^HELP_DIR = .*$", "HELP_DIR = Path('.')", service_src)
    slim = {"title": help_data["title"], "chapters": help_data["chapters"], "intents": help_data["intents"], "faq": help_data["faq"], "context": help_data["context"],
            "topics": {k: {f: v[f] for f in ("id", "title", "number", "chapter", "chapter_title", "screen", "keywords", "related", "summary", "steps", "headings", "text")}
                       for k, v in help_data["topics"].items()}}
    code = (
        "# Search this guide in plain words. Run this cell once, then use ask(\"...\") in any cell.\n"
        "# (The search engine below is the same one BillDesk's own Help screen (F9) uses.)\n"
        + service_src + "\n\n"
        "import json\n"
        f"_DATA = json.loads({json.dumps(json.dumps(slim, ensure_ascii=False), ensure_ascii=False)})\n"
        "_help = HelpService(data=_DATA)\n\n"
        "def ask(question, how_many=5):\n"
        "    \"\"\"Print the best topics for a question, with the first steps.\"\"\"\n"
        "    from IPython.display import Markdown, display\n"
        "    hits = _help.search(question, limit=how_many)\n"
        "    if not hits:\n"
        "        display(Markdown('Nothing matched. Try fewer or simpler words, e.g. **return**, **backup**, **date**.\\n\\nIdeas: ' + '; '.join(_help.suggestions(6))))\n"
        "        return\n"
        "    out = [f'### Results for \"{question}\"']\n"
        "    for n, h in enumerate(hits, 1):\n"
        "        t = h.topic\n"
        "        out.append(f'**{n}. [{t[\"number\"]} {t[\"title\"]}](#{t[\"id\"]})**  _{t[\"chapter_title\"]}_')\n"
        "        if h.why: out.append(f'> {h.why}')\n"
        "        if h.snippet: out.append(h.snippet)\n"
        "        if n <= 3 and h.steps: out.append('Quick steps:\\n' + '\\n'.join(f'{i}. {s}' for i, s in enumerate(h.steps[:4], 1)))\n"
        "    display(Markdown('\\n\\n'.join(out)))\n\n"
        "ask('how do I return goods')\n"
    )
    cells.append({"cell_type": "code", "metadata": {}, "execution_count": None, "outputs": [], "source": code.splitlines(keepends=True)})
    for n, c in enumerate(cells):
        c["id"] = f"guide{n:03d}"
    nb = {"cells": cells, "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
                                       "language_info": {"name": "python"}, "title": guide.title}, "nbformat": 4, "nbformat_minor": 5}
    STEM.with_suffix(".ipynb").write_text(json.dumps(nb, ensure_ascii=False, indent=1), encoding="utf-8")


def main() -> int:
    guide = parse_source(SRC.read_text(encoding="utf-8"))
    validate(guide)
    print(f"{len(guide.chapters)} chapters, {len(guide.topics())} topics, {sum(1 for t in guide.topics() for b in t.blocks if b['t'] == 'image')} screenshots")
    STEM.with_suffix(".md").write_text(published_markdown(guide), encoding="utf-8")
    print("wrote", STEM.name + ".md")
    data = build_help(guide)
    print("wrote app/assets/help/guide.json:", len(data["topics"]), "topics,", len(data["intents"]), "'I want to' phrases,", len(data["faq"]), "questions")
    build_docx(guide)
    print("wrote", STEM.name + ".docx")
    build_pdf(guide)
    print("wrote", STEM.name + ".pdf")
    build_ipynb(guide, data)
    print("wrote", STEM.name + ".ipynb")
    return 0


if __name__ == "__main__":
    sys.exit(main())
